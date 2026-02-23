"""Extract text content from PowerPoint files."""

import os
import subprocess
import tempfile

from pptx import Presentation
from pptx.util import Inches


def convert_ppt_to_pptx(ppt_path: str) -> str:
    """Convert a .ppt file to .pptx using LibreOffice. Returns path to the new file."""
    import shutil

    if shutil.which("libreoffice") is None:
        raise RuntimeError(
            "LibreOffice is not installed. Please upload a .pptx file instead, "
            "or install LibreOffice to enable .ppt support."
        )

    out_dir = tempfile.mkdtemp()
    result = subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pptx", "--outdir", out_dir, ppt_path],
        capture_output=True, text=True, timeout=120,
    )
    if result.returncode != 0:
        raise RuntimeError(f"Failed to convert .ppt file: {result.stderr.strip()}")

    basename = os.path.splitext(os.path.basename(ppt_path))[0] + ".pptx"
    converted = os.path.join(out_dir, basename)
    if not os.path.exists(converted):
        raise RuntimeError("LibreOffice conversion produced no output file.")
    return converted


def extract_slides(pptx_path: str) -> list[dict]:
    """Parse a .pptx file and return structured slide content.

    Returns a list of dicts with keys: slide_number, title, body.
    """
    prs = Presentation(pptx_path)
    slides = []

    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        body_parts = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if shape == slide.shapes.title:
                    title = text
                else:
                    body_parts.append(text)

            if shape.has_table:
                table = shape.table
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(" | ".join(cells))
                body_parts.append("\n".join(rows))

        body = "\n".join(body_parts)

        if title or body:
            slides.append({
                "slide_number": i,
                "title": title,
                "body": body,
            })

    return slides


def format_slides_for_prompt(slides: list[dict]) -> str:
    """Format extracted slides into a text representation for the AI prompt."""
    parts = []
    for slide in slides:
        header = f"--- Slide {slide['slide_number']}"
        if slide["title"]:
            header += f": {slide['title']}"
        header += " ---"
        parts.append(header)
        if slide["body"]:
            parts.append(slide["body"])
        parts.append("")
    return "\n".join(parts)
